from __future__ import annotations

import csv
from collections import defaultdict
from collections.abc import Iterable, Mapping
from pathlib import Path

from openpyxl import load_workbook

from route_builder.models import DayRoute, POI, POIType, SegmentMode, Waypoint

REQUIRED = ("Route ID", "Day", "Sequence", "Name", "Latitude", "Longitude")
POI_REQUIRED = ("Route ID", "Type", "Name", "Latitude", "Longitude")


def _slug(value: str) -> str:
    return "_".join(value.strip().replace("/", " ").split())


def _normalise_mode(value: object) -> SegmentMode:
    raw = str(value or "route").strip().lower().replace("_", "-")
    aliases = {"routed": "route", "offroad": "off-road", "walk": "walking"}
    return SegmentMode(aliases.get(raw, raw))


def _parse_records(records: Iterable[Mapping[str, object]]) -> list[DayRoute]:
    grouped: dict[tuple[str, int], list[Waypoint]] = defaultdict(list)
    for row_number, record in enumerate(records, start=2):
        if not any(value not in (None, "") for value in record.values()):
            continue
        missing = [column for column in REQUIRED if record.get(column) in (None, "")]
        if missing:
            raise ValueError(f"Row {row_number}: missing {', '.join(missing)}")
        point = Waypoint(
            route_id=str(record["Route ID"]).strip(),
            day=int(record["Day"]),
            sequence=int(record["Sequence"]),
            name=str(record["Name"]).strip(),
            latitude=float(record["Latitude"]),
            longitude=float(record["Longitude"]),
            segment_mode=_normalise_mode(record.get("Segment Mode")),
        )
        grouped[(point.route_id, point.day)].append(point)

    days: list[DayRoute] = []
    for (route_id, day), points in sorted(grouped.items()):
        points.sort(key=lambda item: item.sequence)
        sequences = [point.sequence for point in points]
        if len(sequences) != len(set(sequences)):
            raise ValueError(f"Duplicate sequence numbers for {route_id} day {day}")
        days.append(
            DayRoute(
                route_id=route_id,
                day=day,
                name=f"D{day:02d}_{_slug(points[0].name)}_to_{_slug(points[-1].name)}",
                waypoints=points,
            )
        )
    return days


def _parse_poi_records(records: Iterable[Mapping[str, object]]) -> list[POI]:
    pois: list[POI] = []
    for row_number, record in enumerate(records, start=2):
        if not any(value not in (None, "") for value in record.values()):
            continue
        missing = [column for column in POI_REQUIRED if record.get(column) in (None, "")]
        if missing:
            raise ValueError(f"POI row {row_number}: missing {', '.join(missing)}")
        raw_type = str(record["Type"]).strip().lower().replace(" ", "-")
        try:
            poi_type = POIType(raw_type)
        except ValueError:
            poi_type = POIType.OTHER
        day_value = record.get("Day")
        pois.append(
            POI(
                route_id=str(record["Route ID"]).strip(),
                name=str(record["Name"]).strip(),
                poi_type=poi_type,
                latitude=float(record["Latitude"]),
                longitude=float(record["Longitude"]),
                day=int(day_value) if day_value not in (None, "") else None,
                notes=str(record.get("Notes") or "").strip() or None,
            )
        )
    return pois


def parse_csv(path: Path) -> list[DayRoute]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        missing = [column for column in REQUIRED if column not in (reader.fieldnames or [])]
        if missing:
            raise ValueError(f"Missing columns: {', '.join(missing)}")
        return _parse_records(reader)


def parse_excel(path: Path) -> list[DayRoute]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    if "Route Waypoints" not in workbook.sheetnames:
        raise ValueError("Workbook must contain a 'Route Waypoints' sheet")
    rows = workbook["Route Waypoints"].iter_rows(values_only=True)
    try:
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    except StopIteration:
        return []
    missing = [column for column in REQUIRED if column not in headers]
    if missing:
        raise ValueError(f"Missing columns in Route Waypoints: {', '.join(missing)}")
    return _parse_records(dict(zip(headers, row, strict=False)) for row in rows)


def parse_pois(path: Path) -> list[POI]:
    if path.suffix.lower() == ".csv":
        companion = path.with_name(f"{path.stem}_pois.csv")
        if not companion.exists():
            return []
        with companion.open(newline="", encoding="utf-8-sig") as handle:
            reader = csv.DictReader(handle)
            return _parse_poi_records(reader)

    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        workbook = load_workbook(path, data_only=True, read_only=True)
        if "POIs" not in workbook.sheetnames:
            return []
        rows = workbook["POIs"].iter_rows(values_only=True)
        try:
            headers = [str(value).strip() if value is not None else "" for value in next(rows)]
        except StopIteration:
            return []
        return _parse_poi_records(dict(zip(headers, row, strict=False)) for row in rows)

    return []


def parse_input(path: Path) -> list[DayRoute]:
    if path.suffix.lower() == ".csv":
        return parse_csv(path)
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        return parse_excel(path)
    raise ValueError("Unsupported input. Use CSV or XLSX.")
